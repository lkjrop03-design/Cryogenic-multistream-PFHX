"""Results and reporting: turns a converged BVP solution into a per-node
results table, checks it physically (energy balance, limits), and writes
batches of runs to an Excel workbook with a filterable summary sheet.
"""


import os
import numpy as np
import pandas as pd
from hydrogen_pfhx import bvp_model, ortho_para_dynamics


def post_process(solution, streams, reactor, catalyst, topology, boundary_properties):
    """Results table (one row per mesh node) of a converged solution: state
    variables in physical units, transport properties, overall U and heat duty
    of every interface, and the equilibrium para fraction of every catalysed
    stream. Re-evaluates the model at each node, exactly as the solver did.
    """
    z = solution.x
    n = len(z)

    # normalised variants solve dimensionless states: convert back to K and kPa
    if hasattr(bvp_model, 'get_state_scales'):
        scales = bvp_model.get_state_scales(topology, boundary_properties)
    else:
        scales = np.ones(solution.y.shape[0])
    y_physical = solution.y * np.asarray(scales)[:, None]

    columns = {'Z (m)': z}
    # per-node storage for every stream and interface
    per_node = {}

    for label in topology.labels:
        per_node[('h', label)] = np.zeros(n)
        per_node[('velocity', label)] = np.zeros(n)
        per_node[('reynolds', label)] = np.zeros(n)
        per_node[('enthalpy', label)] = np.zeros(n)
        per_node[('density', label)] = np.zeros(n)
        per_node[('cp', label)] = np.zeros(n)
    for pair in topology.interfaces:
        per_node[('U', pair)] = np.zeros(n)
        per_node[('Q', pair)] = np.zeros(n)

    xp_equil = {l: np.zeros(n) for l in topology.packed_labels()}

    for i in range(n):
        # re-evaluate the model at node i exactly as the solver did
        streams = bvp_model.update_parameters(
            solution.y[:, i], streams, topology, boundary_properties)
        temperatures = bvp_model.evaluate_transport_behaviour(
            reactor, catalyst, streams, topology)
        duties, _ = reactor.calculate_heat_transfer_duty(temperatures)

        for label in topology.labels:
            s = streams[label]
            per_node[('h', label)][i] = reactor.h[label]
            per_node[('velocity', label)][i] = s.velocity
            per_node[('reynolds', label)][i] = s.reynolds_number
            per_node[('enthalpy', label)][i] = s.enthalpy
            per_node[('density', label)][i] = s.mass_density
            per_node[('cp', label)][i] = s.specific_heat_capacity

        for pair in topology.interfaces:
            a, b = pair
            # overall U per contact area and heat duty of each interface
            wall = (temperatures[a] + temperatures[b]) / 2
            u, _ = reactor.overall_heat_transfer_coefficient(a, b, wall)
            per_node[('U', pair)][i] = u
            per_node[('Q', pair)][i] = duties[pair]

        for label in topology.packed_labels():
            # equilibrium para fraction at each catalysed stream's own temperature
            xp_eq, _ = ortho_para_dynamics.para_ortho_equilibrium(streams[label].temperature)
            xp_equil[label][i] = xp_eq

    for label in topology.labels:
        e = topology.index[label]
        # state variables
        columns['%s pressure (kPa)' % label] = y_physical[e['P'], :]
        columns['%s temperature (K)' % label] = y_physical[e['T'], :]
        if e['xp'] is not None:
            columns['%s para fraction (mol/mol)' % label] = y_physical[e['xp'], :]

    packed = topology.packed_labels()
    for label in packed:
        # one equilibrium column per catalysed stream; the unlabelled column repeats the first
        columns['%s equilibrium para fraction (mol/mol)' % label] = xp_equil[label]
    if packed:
        columns['Equilibrium para fraction (mol/mol)'] = xp_equil[packed[0]]

    for label in topology.labels:
        # transport properties per stream
        columns['%s velocity (m/s)' % label] = per_node[('velocity', label)]
        columns['%s Reynolds number' % label] = per_node[('reynolds', label)]
        columns['%s film coefficient (W/m2K)' % label] = per_node[('h', label)]
        columns['%s density (kg/m3)' % label] = per_node[('density', label)]
        columns['%s cp (J/kgK)' % label] = per_node[('cp', label)]
        columns['%s enthalpy (J/kg)' % label] = per_node[('enthalpy', label)]

    for pair in topology.interfaces:
        a, b = pair
        # per interface: overall U (per contact area) and heat duty (W/m)
        columns['Overall U %s-%s (W/m2K)' % (a, b)] = per_node[('U', pair)]
        columns['Heat duty %s-%s (W/m)' % (a, b)] = per_node[('Q', pair)]

    return pd.DataFrame(columns)


def check_physical_consistency(results, streams, topology):
    """Per-stream duties and temperature ranges, the overall energy imbalance
    (sum of all duties relative to the largest), range checks (triple point,
    kinetics validity, para fraction within [0, 1]), the catalysed streams'
    approach to equilibrium, and the mesh spacing. Merged into the run
    diagnostics by model.model.
    """
    diagnostics = {}
    total_gained = 0.0

    for label in topology.labels:
        h = results['%s enthalpy (J/kg)' % label].values
        mdot = streams[label].mass_flow_rate
        if topology.boundary_side(label) == 'ya':
            # duty = enthalpy gained between the stream's own inlet and outlet
            duty = (h[-1] - h[0]) * mdot
        else:
            duty = (h[0] - h[-1]) * mdot
        diagnostics['%s_duty_W' % label] = duty
        total_gained += duty

        T = results['%s temperature (K)' % label].values
        # temperature range and solid-phase check
        diagnostics['%s_T_min_K' % label] = float(np.min(T))
        diagnostics['%s_T_max_K' % label] = float(np.max(T))
        diagnostics['%s_T_below_triple' % label] = bool(
            np.any(T < streams[label].triple_point_K))
        if topology.is_packed(label):
            # molar density never below the range of the kinetics fit
            c = results['%s density (kg/m3)' % label].values / streams[label].molecular_mass
            diagnostics['%s_kinetics_below_valid_density' % label] = bool(
                np.any(c < ortho_para_dynamics.KINETICS_MIN_MOLAR_DENSITY))

    # no heat loss is modelled, so all duties should sum to zero
    largest = max((abs(diagnostics['%s_duty_W' % l]) for l in topology.labels),
                  default=1e-9)
    diagnostics['energy_imbalance_pct'] = 100 * abs(total_gained) / max(largest, 1e-9)

    # catalysed streams: inlet/outlet para fraction and approach to equilibrium;
    # unprefixed keys repeat the first catalysed stream
    for k, label in enumerate(topology.packed_labels()):
        xp = results['%s para fraction (mol/mol)' % label].values
        xp_eq = results['%s equilibrium para fraction (mol/mol)' % label].values
        outlet_i = -1 if topology.boundary_side(label) == 'ya' else 0
        inlet_i = 0 if outlet_i == -1 else -1
        values = {'xp_inlet': float(xp[inlet_i]), 'xp_outlet': float(xp[outlet_i]),
                  'xp_equil_outlet': float(xp_eq[outlet_i]),
                  'approach_to_equilibrium': float(xp[outlet_i] - xp_eq[outlet_i]),
                  'xp_bound_violation': bool(np.any(xp < 0) or np.any(xp > 1))}
        for key, val in values.items():
            diagnostics['%s_%s' % (label, key)] = val
        if k == 0:
            diagnostics['xp_reactant_inlet'] = values['xp_inlet']
            diagnostics['xp_reactant_outlet'] = values['xp_outlet']
            diagnostics['xp_equil_outlet'] = values['xp_equil_outlet']
            diagnostics['approach_to_equilibrium'] = values['approach_to_equilibrium']
            diagnostics['xp_bound_violation'] = values['xp_bound_violation']

    # smallest and largest mesh spacing of the converged solution
    dz = np.diff(results['Z (m)'].values)
    diagnostics['min_node_spacing_m'] = float(np.min(dz)) if len(dz) else np.nan
    diagnostics['max_node_spacing_m'] = float(np.max(dz)) if len(dz) else np.nan
    return diagnostics


_INVALID_SHEET_CHARS = set(':\\/?*[]')


def _sanitize_sheet_name(name, used_names):
    """Excel-safe, unique sheet name: max 31 characters, no characters Excel forbids."""
    clean = ''.join(c for c in str(name) if c not in _INVALID_SHEET_CHARS).strip()
    clean = clean[:31] or 'Sheet'
    base, suffix = clean, 1
    while clean in used_names:
        tag = '_%d' % suffix
        clean = base[:31 - len(tag)] + tag
        suffix += 1
    used_names.add(clean)
    return clean


# row colours of the summary sheet per status
_STATUS_FILLS = {
    'OK': 'C6EFCE',
    'CROSSOVER': 'FFC7CE',
    'NOT CONVERGED': 'FFEB9C',
    'ERROR': 'D9D9D9',
}


def _run_status(run):
    """One-word status of a run for the summary: ERROR, NOT CONVERGED,
    CROSSOVER (composite curves cross) or OK. A run that did not converge is
    never reported as CROSSOVER, since its composite curves are meaningless.
    """
    diag = run.get('run_diagnostics') or {}
    if run.get('error') or run.get('results') is None:
        return 'ERROR'
    if not diag.get('success', False):
        return 'NOT CONVERGED'
    if diag.get('composite_feasible') is False:
        return 'CROSSOVER'
    return 'OK'


def _format_summary(ws, summary):
    """Summary sheet layout: filter dropdowns, frozen header, rows coloured by status."""
    from openpyxl.styles import PatternFill
    from openpyxl.utils import get_column_letter

    ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes = 'B2'
    for col_idx, name in enumerate(summary.columns, start=1):
        width = max(len(str(name)), 10) + 2
        ws.column_dimensions[get_column_letter(col_idx)].width = min(width, 40)
    for row_idx, status in enumerate(summary['status'], start=2):
        fill = PatternFill('solid', fgColor=_STATUS_FILLS.get(status, 'FFFFFF'))
        for col_idx in range(1, len(summary.columns) + 1):
            ws.cell(row=row_idx, column=col_idx).fill = fill


def save_results_excel(runs, file_path='output/batch_results.xlsx'):
    """Write a batch of runs to one workbook: a summary sheet (one row per run:
    status, convergence, energy balance, pinch, duties, starting guess, ...),
    then per run its results table and composite-curve table. `runs` is a list
    of dicts with label, results, run_diagnostics and error.
    """
    directory = os.path.dirname(file_path)
    if directory and not os.path.isdir(directory):
        os.makedirs(directory, exist_ok=True)

    summary_rows = []
    for run in runs:
        diag = run.get('run_diagnostics') or {}
        row = {'label': run['label'], 'status': _run_status(run),
               'error': run.get('error')}
        # summary columns, plus every stream's duty
        for key in ('success', 'topology', 'n_states', 'n_iterations', 'n_nodes_final',
                    'runtime_s', 'cpu_time_s', 'guess_runtime_s', 'guess_cpu_time_s',
                    'total_runtime_s', 'total_cpu_time_s', 'energy_imbalance_pct', 'xp_reactant_outlet',
                    'xp_equil_outlet', 'approach_to_equilibrium',
                    'composite_feasible', 'pinch_dT_K', 'pinch_Q_W',
                    'pinch_T_hot_K', 'pinch_T_cold_K',
                    'composite_hot_streams', 'composite_cold_streams',
                    'max_rewarming_excursion_K', 'continuation_steps',
                    'starting_guess', 'xp_bound_violation'):
            row[key] = diag.get(key)
        for key, value in diag.items():
            if key.endswith('_duty_W'):
                row[key] = value
        summary_rows.append(row)

    # summary sheet first, then per run its results and composite-curve sheets
    with pd.ExcelWriter(file_path, engine='openpyxl') as writer:
        used = set()
        summary_sheet = _sanitize_sheet_name('summary', used)
        summary = pd.DataFrame(summary_rows)
        summary.to_excel(writer, sheet_name=summary_sheet, index=False)
        _format_summary(writer.sheets[summary_sheet], summary)
        for run in runs:
            sheet = _sanitize_sheet_name(run['label'], used)
            if run.get('results') is not None:
                run['results'].to_excel(writer, sheet_name=sheet, index=False)
            else:
                pd.DataFrame({'error': [run.get('error') or 'run failed']}).to_excel(
                    writer, sheet_name=sheet, index=False)
            composite = (run.get('run_diagnostics') or {}).get('composite_table')
            if composite is not None:
                comp_sheet = _sanitize_sheet_name(str(run['label'])[:26] + ' comp', used)
                composite.to_excel(writer, sheet_name=comp_sheet, index=False)
